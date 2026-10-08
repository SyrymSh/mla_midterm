"""Build an editable 9-slide PowerPoint using real results when available.

No numerical metrics are fabricated if collection/modeling has not executed.
"""
from __future__ import annotations
import json
from pathlib import Path
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from common import ROOT, CLEAN, RAW, OUT, utc_now

SLIDES = ROOT / 'slides' / 'Video_Game_Midterm_Defense.pptx'
BG = '111923'
PANEL = '1D2832'
WHITE = 'F2F6F4'
MUTED = 'A9B8BD'
GREEN = 'B1ED8C'
AQUA = '63D3C0'
LAVENDER = 'CAB8F0'
AMBER = 'E9BD7B'
DARK = '111923'


def color(code):
    return RGBColor.from_string(code.lstrip('#'))


def rect(slide, x, y, w, h, fill=PANEL, line=None, radius=False):
    typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    obj = slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = color(fill)
    obj.line.fill.background() if line is None else None
    if line:
        obj.line.color.rgb = color(line)
    if radius:
        try:
            obj.adjustments[0] = .12
        except (ValueError, TypeError):
            pass
    return obj


def label(slide, text, x, y, w, h, size=18, bold=False, shade=WHITE,
          font='Aptos', align=None, spacing=1.05, valign=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(.04)
    tf.margin_right = Inches(.04)
    tf.margin_top = Inches(.02)
    tf.margin_bottom = Inches(.02)
    tf.vertical_anchor = valign
    for ii, line in enumerate(str(text).split('\n')):
        p = tf.paragraphs[0] if ii == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(shade)
        p.space_after = Pt(5)
        p.line_spacing = spacing
        if align is not None:
            p.alignment = align
    return tb


def pill(slide, title, x, y, w, fill=GREEN, textshade=DARK):
    rect(slide, x, y, w, .38, fill, radius=True)
    label(slide, title, x+.1, y+.04, w-.2, .31, size=10, bold=True, shade=textshade, align=PP_ALIGN.CENTER)


def base(prs, slide_number, heading, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, BG)
    rect(slide, .6, .57, .06, .6, GREEN)
    if kicker:
        label(slide, kicker.upper(), .84, .48, 11.3, .28, size=10, bold=True, shade=AQUA)
    label(slide, heading, .82, .84, 11.7, .9, size=30, bold=True, shade=WHITE)
    label(slide, 'MACHINE LEARNING ALGORITHMS   /   MIDTERM   /   DIRECTION 42', .72, 7.14, 10.6, .22, size=8, shade=MUTED)
    label(slide, f'{slide_number:02d} / 09', 11.75, 7.11, .9, .25, size=9, bold=True, shade=GREEN, align=PP_ALIGN.RIGHT)
    return slide


def bullet(slide, heading, body, x, y, w, h, accent=GREEN):
    rect(slide, x, y, w, h, PANEL, radius=True)
    rect(slide, x+.21, y+.24, .09, .35, accent, radius=True)
    label(slide, heading, x+.42, y+.21, w-.61, .38, size=15, bold=True)
    label(slide, body, x+.28, y+.7, w-.55, h-.79, size=11, shade=MUTED, spacing=1.15)


def add_image(slide, file, x, y, w, h):
    if not file.exists():
        return False
    from PIL import Image
    with Image.open(file) as image:
        ratio = image.width / image.height
    boxratio = w / h
    if ratio > boxratio:
        display_w = w
        display_h = w / ratio
    else:
        display_h = h
        display_w = h * ratio
    slide.shapes.add_picture(str(file), Inches(x+(w-display_w)/2), Inches(y+(h-display_h)/2),
                             width=Inches(display_w), height=Inches(display_h))
    return True


def fmt(value, f='{:.1f}', fallback='PENDING'):
    try:
        return f.format(float(value)) if pd.notna(value) else fallback
    except (ValueError, TypeError):
        return fallback


def generate():
    reportpath = CLEAN / 'cleaning_report.json'
    report = json.loads(reportpath.read_text(encoding='utf-8')) if reportpath.exists() else {}
    collectionpath = RAW / 'collection_manifest.json'
    collection = json.loads(collectionpath.read_text(encoding='utf-8')) if collectionpath.exists() else {}
    modelpath = OUT / 'model_comparison.csv'
    scores = pd.read_csv(modelpath) if modelpath.exists() else None
    expath = OUT / 'experiment_manifest.json'
    experiment = json.loads(expath.read_text(encoding='utf-8')) if expath.exists() else {}
    measured = scores is not None and len(scores) > 0
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, BG)
    rect(s, 8.65, 0, 4.683, 7.5, PANEL)
    rect(s, .85, 1.16, .1, 1.4, GREEN)
    pill(s, 'DIRECTION 42   •   REGRESSION', 1.15, 1.0, 3.4)
    label(s, 'How long will\nthis game take?', 1.14, 1.83, 7.5, 2.2, size=42, bold=True, spacing=1.02)
    label(s, 'Predicting main-story completion hours from game metadata', 1.17, 4.23, 6.6, .82,
          size=19, shade=MUTED)
    label(s, 'Machine Learning Algorithms   /   Week 5 Midterm', 1.17, 6.73, 8, .32,
          size=12, shade=AQUA)
    for i, (num, title) in enumerate([('01','DATA'),('02','MODELS'),('03','EVALUATION')]):
        y = 1.5 + i * 1.65
        rect(s, 9.2, y, 3.52, 1.29, BG, radius=True)
        label(s, num, 9.5, y+.32, .8, .5, size=28, bold=True, shade=GREEN)
        label(s, title, 10.55, y+.47, 1.7, .35, size=14, bold=True)
    label(s, 'Team: add both names  •  Group: add group', 9.11, 6.76, 4.02, .3, size=10, shade=MUTED)

    # 2 question
    s = base(prs, 2, 'What can metadata tell us?', 'Research question and features')
    rect(s, .8, 1.8, 11.75, 1.18, PANEL, radius=True)
    label(s, 'Can structured features predict average main-story completion time (hours) better than a constant baseline?',
          1.13, 2.08, 10.94, .66, size=20, bold=True)
    items = [
        ('GENRES', 'Action, RPG, strategy, adventure', AQUA),
        ('RELEASE YEAR', 'Game generation / release period', GREEN),
        ('USD PRICE', 'Missing for all games in this IGDB sample', LAVENDER),
        ('DEVELOPER TIER', 'Catalog-frequency proxy, not studio size', AMBER),
        ('GAME TAGS', 'No community tags returned in this sample', AQUA),
    ]
    for i, (h,b,c) in enumerate(items):
        x = .82 + (i % 3) * 4.03
        y = 3.3 + (i // 3) * 1.45
        bullet(s,h,b,x,y,3.8,1.24,c)
    pill(s, 'Description text reserved for Final NLP experiment', 9.15, 6.15, 3.42, fill=GREEN)

    # 3 sources
    s = base(prs, 3, 'Real completion time + game metadata', 'Provenance and collection')
    bullet(s, 'DIRECT SOURCE  /  IGDB',
           'Official game_time_to_beats API. hastily = time to credits in seconds; ÷ 3,600 yields hours. Requires Twitch developer credentials.',
           .85, 1.82, 5.66, 2.1, GREEN)
    bullet(s, 'ALTERNATIVE  /  PUBLIC COMPILATION',
           'Game Gauntlets / steamdb: HLTB main-story hours with Steam tags, descriptions, genres and price. Disclose that this is third-party compiled data.',
           6.73, 1.82, 5.72, 2.1, AQUA)
    rect(s, .85, 4.15, 11.6, 1.66, PANEL, radius=True)
    label(s, 'REPRODUCIBLE EVIDENCE', 1.16, 4.42, 4.5, .38, 15, True, GREEN)
    label(s, 'Collection code  •  raw JSONL  •  UTC timestamp  •  SHA256  •  source links  •  cleaning log',
          1.16, 4.98, 10.4, .46, 14)
    label(s, 'Source access & terms must be rechecked; IGDB authentication is required.',
          .93, 6.2, 10.8, .43, 12, False, AMBER)

    # 4 dataset + preparation
    s = base(prs, 4, 'One row = one game with known hours', 'Cleaning and data card')
    data_metrics = [('RAW ROWS', f'{report.get("raw_rows", "PENDING")}', AQUA),
                    ('CLEAN ROWS', f'{report.get("cleaned_rows", "PENDING")}', GREEN),
                    ('MEDIAN HOURS', fmt(report.get('median_target_hours')), AMBER),
                    ('MISSING PRICE', fmt(report.get('missing_price_percent'), '{:.1f}%'), LAVENDER)]
    for i, (txt, val, shade) in enumerate(data_metrics):
        x = .9 + i*3.0
        rect(s, x, 1.86, 2.75, 1.4, PANEL, radius=True)
        label(s, val, x+.18, 2.02, 2.42, .52, 27, True, shade)
        label(s, txt, x+.19, 2.7, 2.39, .34, 10, True, MUTED)
    label(s, 'Collection and cleaning checks', 1.05, 3.64, 7.9, .38, 19, True)
    bullet(s, '01  VALID TARGET', 'Keep numeric, positive main-story hours; never impute the target.', .9, 4.12, 3.78, 1.54, GREEN)
    bullet(s, '02  DE-DUPLICATE', 'Remove repeated IDs and normalized title-year records.', 4.77, 4.12, 3.78, 1.54, AQUA)
    bullet(s, '03  PRESERVE TEXT', 'Store original descriptions for Final; exclude from Midterm regressors.', 8.66, 4.12, 3.78, 1.54, LAVENDER)
    label(s, 'Missing numeric features are imputed after the train/test split, inside each training fold.', 1.05, 6.2, 11.5, .43, 12, False, MUTED)

    # 5 EDA
    s = base(prs, 5, 'Explore duration and genre imbalance', 'Exploratory data analysis')
    figs = OUT / 'figures'
    rect(s, .87, 1.91, 5.74, 4.16, WHITE, radius=True)
    rect(s, 6.73, 1.91, 5.74, 4.16, WHITE, radius=True)
    left = add_image(s, figs / 'target_distribution.png', 1.06, 2.05, 5.34, 3.84)
    right = add_image(s, figs / 'genre_counts.png', 6.94, 2.05, 5.31, 3.84)
    if not left:
        label(s, 'DURATION DISTRIBUTION', 1.2, 3.06, 5.0, .5, 21, True, DARK, align=PP_ALIGN.CENTER)
        label(s, 'Chart created after collection and training', 1.35, 3.64, 4.72, .7, 14, False, '687A83', align=PP_ALIGN.CENTER)
    if not right:
        label(s, 'GENRE COVERAGE', 7.15, 3.06, 4.9, .5, 21, True, DARK, align=PP_ALIGN.CENTER)
        label(s, 'Chart created after collection and training', 7.22, 3.64, 4.78, .7, 14, False, '687A83', align=PP_ALIGN.CENTER)
    label(s, 'Main-story lengths are positive and may be strongly right-skewed; verify the chart before claiming a shape.',
          .98, 6.29, 11.53, .43, 12, False, MUTED)

    # 6 method
    s = base(prs, 6, 'A fair model comparison', 'Methodology & leakage controls')
    rect(s, .87, 1.87, 11.59, 1.32, PANEL, radius=True)
    label(s, '80% TRAIN', 1.24, 2.22, 2.12, .55, 23, True, GREEN)
    label(s, '3-fold CV on train only', 3.6, 2.29, 3.9, .43, 16, True, AQUA)
    label(s, '20% TEST (fixed)', 9.12, 2.24, 2.97, .43, 19, True, AMBER)
    rect(s, 1.15, 2.88, 8.0, .075, GREEN)
    model_items = [
        ('BASELINE', 'Median DummyRegressor', AMBER),
        ('LINEAR', 'Ordinary least squares', AQUA),
        ('REGULARIZED', 'Ridge regression', GREEN),
        ('NONLINEAR', 'Decision tree (limited depth)', LAVENDER),
        ('NEIGHBORHOOD', 'K-nearest neighbors', AQUA),
    ]
    for i,(top,bottom,shade) in enumerate(model_items):
        x = .85 + (i%3)*4.04
        y = 3.53 + (i//3)*1.21
        bullet(s, top, bottom, x, y, 3.78, 1.03, shade)
    label(s, 'Training-only: imputation, scaling, developer tier and multi-tag vocabulary.  Seed = 42.',
          1.0, 6.47, 11.5, .38, 12, False, MUTED)

    # 7 results
    s = base(prs, 7, 'How accurately do we predict hours?', 'Held-out model results')
    rect(s, .85, 1.83, 11.58, 4.62, PANEL, radius=True)
    xs = [1.19, 6.2, 8.02, 9.92, 11.08]
    labels = ['MODEL', 'CV MAE', 'TEST MAE', 'TEST RMSE', 'R²']
    for x, header in zip(xs, labels):
        label(s, header, x, 2.12, 1.4 if x>6 else 4.7, .38, 11, True, GREEN)
    rect(s, 1.12, 2.61, 10.93, .018, '3B4B52')
    order = ['Dummy (median)', 'Linear Regression', 'Ridge Regression', 'Decision Tree', 'K-Nearest Neighbors']
    for i, name in enumerate(order):
        y = 2.82 + i*.57
        label(s, name, 1.2, y, 4.84, .36, 13, True, WHITE)
        r = scores[scores.model == name].iloc[0] if measured and (scores.model == name).any() else None
        vals = (fmt(r.train_cv_MAE_hours, '{:.2f}') if r is not None else '—',
                fmt(r.MAE_hours, '{:.2f}') if r is not None else '—',
                fmt(r.RMSE_hours, '{:.2f}') if r is not None else '—',
                fmt(r.R2, '{:.3f}') if r is not None else '—')
        for x,v in zip(xs[1:],vals):
            label(s,v,x,y,1.45,.36,13,False,AQUA if r is not None else MUTED)
    if measured:
        winner = experiment.get('best_cv_model','See experiment manifest')
        msg = f'Selected by training CV: {winner}  •  Test scores are measured, not used to choose the model.'
        shade = GREEN
    else:
        msg = 'PENDING REAL DATA & EXECUTION — these blank cells are NOT measured results.'
        shade = AMBER
    label(s,msg,1.03,6.62,11.6,.36,12,True,shade)

    # 8 errors
    s = base(prs, 8, 'Where can our estimates fail?', 'Error analysis & limitations')
    rect(s, .84, 1.89, 7.28, 4.37, WHITE, radius=True)
    if not add_image(s, figs / 'error_by_time.png', 1.05, 2.04, 6.84, 4.03):
        label(s, 'ERROR BY GAME LENGTH', 1.24, 3.22, 6.49, .5, 22, True, DARK, align=PP_ALIGN.CENTER)
        label(s, 'Chart pending actual held-out predictions', 1.5, 3.88, 5.95, .43, 14, False, '687A83', align=PP_ALIGN.CENTER)
    bullet(s, 'LONG-TAIL TITLES', 'Very lengthy games can yield large absolute errors, even if relative error is moderate.',
           8.42, 1.93, 4.02, 1.32, AMBER)
    bullet(s, 'METADATA GAPS', 'Prices and community tags are absent for every game, so these fields add no predictive signal.',
           8.42, 3.36, 4.02, 1.32, AQUA)
    bullet(s, 'SOURCE BIAS', 'Completion times come from participating players and are not exact for every person.',
           8.42, 4.82, 4.02, 1.36, LAVENDER)
    label(s, 'Check actual residuals and genre groups before attributing specific errors to any cause.',
          1.02, 6.57, 11.5, .36, 12, False, MUTED)

    # 9 next steps
    s = base(prs, 9, 'From structured features to text', 'Endterm + Final roadmap')
    bullet(s, 'ENDTERM   •   WEEK 9',
           'Tune ensemble models; run K-means segmentation of game profiles; examine cluster meanings; compare a structured neural network.',
           .85, 1.86, 5.63, 2.27, AQUA)
    bullet(s, 'FINAL   •   AFTER WEEK 10',
           'TF-IDF game descriptions. Compare structured-only, description-only and combined feature pipelines on the same fixed split, then build a local demo.',
           6.75, 1.86, 5.72, 2.27, GREEN)
    rect(s, .86, 4.49, 11.61, 1.47, PANEL, radius=True)
    label(s, 'REPORTING & RESPONSIBILITIES', 1.2, 4.75, 10.6, .42, 16, True, WHITE)
    label(s, 'Add both members and actual work split. Disclose AI-assisted scaffolding; run & validate the real collection and models.',
          1.2, 5.29, 10.5, .48, 12, False, MUTED)
    label(s, 'Main deliverables: notebook  •  collection code  •  raw + clean data  •  these slides',
          .99, 6.47, 11.4, .4, 12, True, GREEN)

    SLIDES.parent.mkdir(parents=True, exist_ok=True)
    prs.save(SLIDES)
    print('Saved', SLIDES, '|', len(prs.slides), 'slides | measured results:', measured)
    return SLIDES


if __name__ == '__main__':
    generate()
